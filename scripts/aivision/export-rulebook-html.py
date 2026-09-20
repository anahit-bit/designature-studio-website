#!/usr/bin/env python
r"""
Render AI-Vision-Rulebook.xlsx as one browsable HTML page — a read-only
visual view of the Rules sheet for the owner, who finds Excel easier to scan
but a rendered page easier to open and read than launching Excel each time.

    python scripts/aivision/export-rulebook-html.py

Reads   E:\Business\Claude\Doc\cards\redesign-my-room\rules\AI-Vision-Rulebook.xlsx
Writes  E:\Business\Claude\Doc\cards\redesign-my-room\rules\AI-Vision-Rulebook.html

GENERATED. Edit the workbook, re-run this (and compile-rulebook.py) together —
never hand-edit the HTML. Whenever a Rules-sheet row changes, both
rulebook.generated.ts (via compile-rulebook.py) and this HTML view need
re-running, or the three copies of "what the rules are" drift apart.
"""
import html
import os
import sys
from datetime import datetime

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

RULES_DIR = r"E:\Business\Claude\Doc\cards\redesign-my-room\rules"
XLSX = os.path.join(RULES_DIR, "AI-Vision-Rulebook.xlsx")
OUT = os.path.join(RULES_DIR, "AI-Vision-Rulebook.html")

LEVEL_COLOR = {
    "ABSOLUTE": "#9E5E41",  # terracotta — never bends
    "HARD": "#0047AB",      # cobalt — enforced
    "ADVISORY": "#6B6B6B",
    "WARN": "#B08900",
}


def e(v) -> str:
    return html.escape(str(v if v is not None else ""))


def main() -> None:
    if not os.path.exists(XLSX):
        sys.exit(f"Workbook not found: {XLSX}")
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Rules"]
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header)}

    groups: dict[str, list] = {}
    order: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[col["ID"]]:
            continue
        group = str(row[col["Group"]] or "Ungrouped")
        if group not in groups:
            groups[group] = []
            order.append(group)
        groups[group].append(row)

    rows_html = []
    for group in order:
        rows_html.append(f'<tr class="group-row"><td colspan="6">{e(group)}</td></tr>')
        for row in groups[group]:
            rid = row[col["ID"]]
            rule = row[col["Rule"]]
            level = str(row[col["Level"]] or "")
            status = str(row[col["Status"]] or "")
            evidence = row[col["Why / evidence"]]
            enforced = row[col["Enforced by"]]
            color = LEVEL_COLOR.get(level.upper(), "#6B6B6B")
            status_dim = "" if status.strip().lower() in ("active",) else ' style="opacity:.55"'
            rows_html.append(
                f'<tr{status_dim}>'
                f'<td class="id">{e(rid)}</td>'
                f'<td><span class="level" style="background:{color}">{e(level)}</span></td>'
                f'<td>{e(status)}</td>'
                f'<td class="rule">{e(rule)}</td>'
                f'<td class="ev">{e(evidence)}</td>'
                f'<td class="ev">{e(enforced)}</td>'
                f'</tr>'
            )

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    src_mtime = datetime.fromtimestamp(os.path.getmtime(XLSX)).strftime("%Y-%m-%d %H:%M")

    doc = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>AI Vision Rulebook — Redesign My Room</title>
<style>
  body {{ font-family: -apple-system, Segoe UI, Arial, sans-serif; background:#FBF7EE; color:#0A0A0A; margin:0; padding:32px; }}
  h1 {{ color:#0047AB; margin:0 0 4px; }}
  .meta {{ color:#6B6B6B; font-size:13px; margin-bottom:24px; }}
  table {{ width:100%; border-collapse:collapse; background:#fff; box-shadow:0 1px 3px rgba(0,0,0,.08); }}
  th, td {{ text-align:left; padding:8px 10px; border-bottom:1px solid #DAD2C3; vertical-align:top; font-size:13px; }}
  th {{ background:#0047AB; color:#fff; position:sticky; top:0; }}
  td.id {{ font-weight:600; white-space:nowrap; }}
  td.rule {{ max-width:420px; }}
  td.ev {{ max-width:340px; color:#333; }}
  tr.group-row td {{ background:#F0F4FF; font-weight:700; color:#0047AB; letter-spacing:.03em; text-transform:uppercase; font-size:11px; padding-top:16px; }}
  .level {{ color:#fff; padding:2px 8px; border-radius:3px; font-size:11px; font-weight:600; white-space:nowrap; }}
  .note {{ background:#FFF7E8; border:1px solid #DAD2C3; padding:12px 16px; border-radius:4px; margin-bottom:20px; font-size:13px; }}
</style>
</head>
<body>
<h1>AI Vision Rulebook</h1>
<div class="meta">Generated {generated_at} from AI-Vision-Rulebook.xlsx (workbook last saved {src_mtime}) — read-only view, GENERATED, never hand-edit. Regenerate: <code>python scripts/aivision/export-rulebook-html.py</code></div>
<div class="note">This is a browsing view. The <strong>xlsx is still the editable source of truth</strong> — edit it there, then re-run both <code>compile-rulebook.py</code> (updates the live prompt) and <code>export-rulebook-html.py</code> (updates this page) together, and update <code>rulebook.md</code>'s prose if the change is conceptual, not just wording. Rows shown dimmed have Status other than "Active".</div>
<table>
<thead><tr><th>ID</th><th>Level</th><th>Status</th><th>Rule</th><th>Why / evidence</th><th>Enforced by</th></tr></thead>
<tbody>
{"".join(rows_html)}
</tbody>
</table>
</body>
</html>
"""
    with open(OUT, "w", encoding="utf-8", newline="\n") as f:
        f.write(doc)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()

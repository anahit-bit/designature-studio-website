# -*- coding: utf-8 -*-
"""
Rulebook edit, 2026-09-07 (second) - the ceiling gets counted.

RD5 has been written out in full since 2026-07-13 and is still broken in
essentially every generation: a lit perimeter cove appeared in all six graded
bathroom images, before and after the plumbing work, on rooms whose ceilings are
plainly flat. One also cut a recessed niche into an existing tiled duct.

RD22 is explicit about what to do with a rule that keeps failing after being
stated: escalate from prompt text to verify-and-retry. So this is not a rewording
of RD5 - the wording was never the problem. Two changes:

  1. RD5's prompt text gains the two things it did not name - downlights sunk
     into the plane, and the positive instruction to keep a flat ceiling flat.
  2. RD28 records the enforcement: spatialAnalysis measures the ceiling, the
     survey states it positively, and inventedCeiling() counts relief that was
     not in the source, with one corrective retry.

Requires the workbook to be CLOSED in Excel.

Run once:  python scripts/aivision/apply-owner-edits-2026-09-07b.py
Then:      python scripts/aivision/compile-rulebook.py
"""
import sys
from openpyxl import load_workbook

XLSX = r"E:\Business\Claude\_Plan\Website\AI-Vision-Rulebook.xlsx"

RD5_PROMPT = (
    "Leave the ceiling as the single flat plane it already is, at its original height. Repaint it if the "
    "style calls for it and change the light fittings on it - but add nothing to it: no beams, soffits, "
    "bulkheads, coffers, dropped or tray sections, perimeter coves, shadow gaps, LED channels, floating "
    "panels, or plank, slat or timber cladding, and no downlights sunk INTO the plane. A light fitting "
    "hangs from the ceiling or sits on its surface; it is never recessed into it unless the photograph "
    "already shows recessed lights. If the photograph shows a plain flat ceiling, the renovated room has "
    "a plain flat ceiling - that is the correct result, not an unfinished one."
)

RD28 = {
    "ID": "RD28",
    "Group": "Enforcement",
    "Rule": (
        "The ceiling plane is measured before generation and counted after it. Relief present in the "
        "output but not in the source - cove, dropped section, coffer, beam, or downlights recessed into "
        "the plane - triggers one corrective regeneration."
    ),
    "Level": "HARD",
    "Why / evidence": (
        "RD22 escalation for RD5, which is the owner's most repeated complaint and has survived every "
        "rewording since 2026-07-13. On 2026-09-07 a lit perimeter cove appeared in all six graded "
        "bathroom generations - both engines of the plumbing A/B, on rooms with plainly flat ceilings. "
        "Prompt text does not enforce; a count does, exactly as it did for openings (RD25) and plumbing "
        "(RD27)."
    ),
    "Enforced by": (
        "spatialAnalysis.ceiling {flat, features, cornice} rendered as a CEILING SURVEY line stating the "
        "plane positively; inventedCeiling() diff in imageGeneration.ts with one corrective retry."
    ),
    "Status": "active",
    "Prompt text": "",
    "Engines": "none",
    "Prompt section": "",
    "Owner note": "",
}


def main() -> None:
    try:
        wb = load_workbook(XLSX)
    except PermissionError:
        sys.exit("Workbook is open in Excel. Close it and run this again.")

    ws = wb["Rules"]
    header = [(c.value or "").strip() if isinstance(c.value, str) else "" for c in ws[1]]
    col = {n: i + 1 for i, n in enumerate(header)}
    ids = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=col["ID"]).value
        if v:
            ids[str(v).strip()] = r

    if "RD5" not in ids:
        sys.exit("RD5 not found on the Rules sheet.")
    ws.cell(row=ids["RD5"], column=col["Prompt text"], value=RD5_PROMPT)
    print("Rules: RD5 prompt text now names recessed downlights and states the flat case positively")

    if "RD28" in ids:
        print("Rules: RD28 already present, left alone")
    else:
        r = ws.max_row + 1
        for name, value in RD28.items():
            if name in col:
                ws.cell(row=r, column=col[name], value=value)
        print("Rules: added RD28")

    try:
        wb.save(XLSX)
    except PermissionError:
        sys.exit("Workbook is open in Excel. Close it and run this again.")
    print("\nSaved %s" % XLSX)


if __name__ == "__main__":
    main()

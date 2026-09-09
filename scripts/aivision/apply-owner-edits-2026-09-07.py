# -*- coding: utf-8 -*-
"""
Rulebook edit, 2026-09-07 - plumbing is evidence-bound.

Owner reported a bathroom concept that plumbed a wall-hung toilet and a heated
towel rail into a wall the photograph shows bare. Their reasoning, which is
sharper than the 2026-09-04 rewrite: a toilet needs a soil pipe. If the photo
does not show the fixture, there is no evidence of drainage on that wall, so the
concept is not buildable.

The 2026-09-04 rewrite made the BATHING fixture conditional ("a bath stays a
bath, a shower stays a shower") but left "Include ... a toilet" as an
unconditional MUST, and then repeated it in the fallback clause - so the toilet
was mandated twice. RD24 could not catch it either: its list of things a
programme may not authorise is architectural (opening, wall, recess, level),
and a toilet does not read as building work.

Three changes:
  1. RD27 - plumbed fixtures are fixed fabric. Kept where they are, never added
     to a wall with no drainage evidence.
  2. RD24 gains plumbing and drainage, so the governing rule covers it.
  3. The bathroom programme stops mandating a toilet, and defers to the survey.

Backed by a measured survey: spatialAnalysis now reports `plumbing` (fixtures
plus soil stacks / pipe boxing), rendered into the prompt as a positive
inventory - or as an explicit "this room has NO drainage evidence" when empty.
RD26 adds the photo-specific note, and RD25's post-generation check now counts
plumbed fixtures as well as openings.

Requires the workbook to be CLOSED in Excel.

Run once:  python scripts/aivision/apply-owner-edits-2026-09-07.py
Then:      python scripts/aivision/compile-rulebook.py
"""
import sys
from openpyxl import load_workbook

XLSX = r"E:\Business\Claude\_Plan\Website\AI-Vision-Rulebook.xlsx"

RD27 = {
    "ID": "RD27",
    "Group": "Fixed fabric",
    "Rule": (
        "Plumbed fixtures are fixed fabric, not furnishings. A toilet, bidet, bath, shower, basin or "
        "heated towel rail may appear only where the photograph already shows one, or shows the boxing, "
        "duct or soil stack that serves it. Keep every one the photo shows; never add one to a wall with "
        "no drainage evidence."
    ),
    "Level": "ABSOLUTE",
    "Why / evidence": (
        "2026-09-07: a bathroom photographed with a vanity, window, radiator and door - and no toilet - "
        "came back with a wall-hung toilet and a heated towel rail plumbed into a bare wall. Owner: "
        "'if there is no toilet in the image, then there is no plumbing, then it should not invent it.' "
        "The 2026-09-04 rewrite had made the bathing fixture conditional but left the toilet mandatory "
        "twice over, and RD24's list was architectural so it did not cover fixtures."
    ),
    "Enforced by": (
        "spatialAnalysis.plumbing (measured survey incl. soil stacks) rendered as a PLUMBING SURVEY line; "
        "RD26 photo-specific note; RD25 post-generation count via inventedPlumbing() with one corrective "
        "retry."
    ),
    "Status": "active",
    "Prompt text": (
        "Plumbed fixtures are fixed fabric, not furnishing. A toilet, bidet, bath, shower, basin or heated "
        "towel rail needs a waste pipe and a feed, so one may appear ONLY where the photograph already "
        "shows that fixture, or shows the boxing, duct or soil stack that serves it. Keep every plumbed "
        "fixture the photograph shows, on its own wall at its own size - replace or refinish it in the "
        "target style, but never move it to another wall and never delete it. Never add a plumbed fixture "
        "to a wall the photograph shows bare: floor space is not drainage. A bathroom photographed without "
        "a toilet is rendered without a toilet."
    ),
    "Engines": "gemini",
    "Prompt section": "architecture",
    "Owner note": "",
}

RD24_NEW_PROMPT = (
    "The room programme that follows lists what the finished room CONTAINS - furniture, fittings, "
    "lighting and styling. It can never authorise building work. If any part of it would need an "
    "opening, doorway, arch, staircase, window, wall, recess, change of level, or a NEW WATER SUPPLY "
    "OR WASTE PIPE that this photograph does not already show, DROP that part and furnish what is "
    "actually there. A room that reads as a dead end, or as small, or as awkwardly shaped, or as "
    "lacking a fixture you would expect it to have, is the correct answer when that is the room in "
    "the photograph."
)

# The toilet stops being mandatory and defers to the measured survey.
BATHROOM = (
    "The room MUST be a fully realized BATHROOM. It MUST be properly TILED: real wall tile behind the "
    "vanity and through the shower or bath area, and tiled or stone flooring, with visible grout lines "
    "and a tile format, colour and laying pattern chosen to suit the target style. Tile is a finish laid "
    "flat onto the walls that are already there - it never changes their plane. THE PLUMBING SURVEY "
    "ABOVE DECIDES WHICH FIXTURES EXIST, AND THIS PROGRAMME CANNOT ADD TO IT: keep every toilet, bidet, "
    "bath, shower and basin the photograph shows, each on its own wall at its own size, and add NONE "
    "that it does not. A bathroom photographed without a toilet is rendered without a toilet, and a wall "
    "shown bare stays bare - that is a correct result, not an incomplete one. Style what is genuinely "
    "there: replace or refinish the fixtures, the vanity, the mirror, the tiling, the lighting and the "
    "floor, and add the unplumbed things freely - mirror, storage, towel bars or rings with real towels, "
    "sconces or vanity lighting, a bath mat, plants and styling. Do NOT carve a new recess, niche, "
    "wet-room enclosure or partition to house a fixture, do NOT enlarge the room to fit one, and do NOT "
    "add a heated towel rail unless the photograph already has one. Do NOT leave the walls as bare "
    "plaster or paint alone, and do NOT include living-room furniture, bedroom furniture, or dining "
    "tables. Every fixture must be a real bathroom fixture."
)


def main() -> None:
    try:
        wb = load_workbook(XLSX)
    except PermissionError:
        sys.exit("Workbook is open in Excel. Close it and run this again.")

    ws = wb["Rules"]
    header = [(c.value or "").strip() if isinstance(c.value, str) else "" for c in ws[1]]
    col = {n: i + 1 for i, n in enumerate(header)}
    for need in ("ID", "Prompt text", "Prompt section", "Engines"):
        if need not in col:
            sys.exit("Rules sheet is missing the '%s' column." % need)

    ids = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=col["ID"]).value
        if v:
            ids[str(v).strip()] = r

    if "RD24" in ids:
        ws.cell(row=ids["RD24"], column=col["Prompt text"], value=RD24_NEW_PROMPT)
        print("Rules: RD24 prompt text now covers water supply and waste")
    else:
        sys.exit("RD24 not found - run apply-owner-edits-2026-09-04.py first.")

    if "RD27" in ids:
        print("Rules: RD27 already present, left alone")
    else:
        r = ws.max_row + 1
        for name, value in RD27.items():
            if name in col:
                ws.cell(row=r, column=col[name], value=value)
        print("Rules: added RD27")

    ws = wb["Room Programs"]
    header = [(c.value or "").strip() if isinstance(c.value, str) else "" for c in ws[1]]
    col = {n: i + 1 for i, n in enumerate(header)}
    hit = False
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=col["Room key"]).value or "").strip() == "bathroom":
            ws.cell(row=r, column=col["Programme rule handed to the model"], value=BATHROOM)
            hit = True
            print("Room Programs: bathroom rewritten (%d chars)" % len(BATHROOM))
    if not hit:
        sys.exit("No 'bathroom' row on the Room Programs sheet.")

    try:
        wb.save(XLSX)
    except PermissionError:
        sys.exit("Workbook is open in Excel. Close it and run this again.")
    print("\nSaved %s" % XLSX)


if __name__ == "__main__":
    main()

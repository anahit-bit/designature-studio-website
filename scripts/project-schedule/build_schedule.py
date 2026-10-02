"""Build the Designature Studio project schedule workbook.

Run:  python3 scripts/project-schedule/build_schedule.py
Out:  docs/scheduling/DesignatureStudio-Project-Schedule.xlsx

Change PHASES, PROJECTS or DEFAULTS below and rerun to rebuild, or edit the
workbook directly: durations, dates, bars and the workload row are all formulas
driven by the Settings sheet and the blue input cells.
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "docs/scheduling/DesignatureStudio-Project-Schedule.xlsx"

# Reference values from Lara: one project at a time, 6 h/day, Mon to Fri, 50 m2.
DEFAULTS = dict(hours_per_day=6, days_per_week=5, ref_area=50, gap_days=0)

# (phase, owner, reference weeks at 50 m2, scaling exponent, bar colour, reasoning)
# Lara's four phases are her numbers. Everything owned by Anahit is a tracking
# placeholder with invented durations. Weeks = 0 skips a phase.
PHASES = [
    ("Brief and concept", "Anahit", 1, 0.0, "DADADA", "Fixed. A conversation and a mood board take the same time at any size. Placeholder."),
    ("Floor planning", "Lara", 1, 0.5, "9DB4A0", "Layout effort follows room count and circulation problems, not area. Weak scaling."),
    ("3D modeling", "Lara", 2, 0.9, "C9A27E", "Every surface gets modeled and materials assigned, so it tracks area closely."),
    ("Client approval", "Anahit", 1, 0.0, "C4C4C4", "Waiting on the client. Fixed. Placeholder."),
    ("Technical drawings", "Lara", 1, 0.8, "8FA7C2", "Sheet count follows rooms and wall and ceiling detail. Mostly area driven."),
    ("Furniture drawings", "Lara", 2, 0.7, "B58E9E", "Driven by the number of custom pieces, which grows slower than area."),
    ("Procurement", "Anahit", 3, 0.0, "E8D9B5", "Ordering and lead times, mostly fixed. Placeholder."),
    ("Site supervision", "Anahit", 6, 0.5, "E0C3B0", "Longer for bigger sites but not in proportion. Placeholder."),
]

# PLACEHOLDER projects. Replace with the real ones.
# (name, client, area m2, complexity, share of Lara's time, start date, status, note)
PROJECTS = [
    ("Example 01", "Client A", 50, 1.0, 1.0, date(2026, 10, 5), "Active", "Reference size, Lara full time"),
    ("Example 02", "Client B", 72, 1.0, 0.5, date(2026, 10, 19), "Planned", "Parallel with 01 and 03, half her time"),
    ("Example 03", "Client C", 38, 0.9, 0.5, date(2026, 10, 26), "Planned", "Small, simple layout"),
    ("Example 04", "Client D", 95, 1.2, 1.0, date(2026, 11, 16), "Planned", "Lots of custom joinery"),
    ("Example 05", "Client E", 60, 1.0, 0.6, date(2026, 12, 7), "Planned", ""),
    ("Example 06", "Client F", 120, 1.1, 0.4, date(2027, 1, 11), "Planned", "House, runs in the background"),
]
MAX_ROWS = 30
GANTT_WEEKS = 60

FONT = "Arial"
def f(**k):
    return Font(name=FONT, size=k.pop("size", 10), **k)
BLUE, BLACK, BOLD = f(color="0000FF"), f(), f(bold=True)
GREY = f(italic=True, color="666666")
WHITE_B = f(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2B2B2B")
INPUT_FILL = PatternFill("solid", fgColor="FFF9D6")
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
MID = Alignment(horizontal="center")
DATE_FMT = "dd mmm yy;;"          # zero (empty row) shows blank
NP = len(PHASES)

wb = Workbook()

# ============================ Settings ============================
s = wb.active
s.title = "Settings"
s.sheet_view.showGridLines = False
s["B1"] = "Settings and reference durations"
s["B1"].font = f(bold=True, size=14)
s["B2"] = "Blue text on cream is an input. Everything else is a formula."
s["B2"].font = GREY

G_ROWS = 18 + 2 * NP  # not used for layout, kept simple
top = [
    (4, "Working hours per day (Lara)", DEFAULTS["hours_per_day"], "Her average, from Lara."),
    (5, "Working days per week", DEFAULTS["days_per_week"], "Monday to Friday, confirmed."),
    (6, "Reference apartment area (m2)", DEFAULTS["ref_area"], "Size the reference weeks apply to."),
    (7, "Idle working days between phases", DEFAULTS["gap_days"], "Extra waiting time between any two phases. 0 = back to back."),
]
for r, label, val, note in top:
    s.cell(r, 2, label).font = BLACK
    c = s.cell(r, 3, val)
    c.font, c.fill, c.border, c.alignment = BLUE, INPUT_FILL, BORDER, MID
    s.cell(r, 4, note).font = GREY
s["B8"] = "Gantt first week (Monday)"
s["C8"] = f"=MIN(Schedule!$G$5:$G${4 + MAX_ROWS})-WEEKDAY(MIN(Schedule!$G$5:$G${4 + MAX_ROWS}),2)+1"
s["C8"].number_format, s["C8"].font, s["C8"].alignment = "dd mmm yyyy", BLACK, MID
s["D8"] = "Monday of the earliest project start. Overwrite with a date to pin the timeline."
s["D8"].font = GREY

PH0 = 12
for i, h in enumerate(["Phase", "Owner", "Reference weeks", "Scaling exponent", "Ref. working days", "Ref. hours", "Why this scaling"]):
    c = s.cell(PH0 - 1, 2 + i, h)
    c.font, c.fill, c.alignment, c.border = WHITE_B, HEAD_FILL, CENTER, BORDER
s.row_dimensions[PH0 - 1].height = 30
for i, (name, owner, wk, ex, _col, why) in enumerate(PHASES):
    r = PH0 + i
    s.cell(r, 2, name).font = BLACK
    for col, v in ((3, owner), (4, wk), (5, ex)):
        c = s.cell(r, col, v)
        c.font, c.fill, c.alignment = BLUE, INPUT_FILL, MID
    s.cell(r, 6, f'=D{r}*$C$5').font = BLACK
    s.cell(r, 7, f'=IF(C{r}="Lara",F{r}*$C$4,0)').font = BLACK
    s.cell(r, 8, why).font = f(color="444444", size=9)
    for col in range(2, 9):
        s.cell(r, col).border = BORDER
    s.cell(r, 6).alignment = s.cell(r, 7).alignment = MID
PHL = PH0 + NP - 1
tr = PHL + 1
s.cell(tr, 2, "Total").font = BOLD
for col, fm in ((4, f"=SUM(D{PH0}:D{PHL})"), (6, f"=SUM(F{PH0}:F{PHL})"), (7, f"=SUM(G{PH0}:G{PHL})")):
    c = s.cell(tr, col, fm)
    c.font, c.alignment = BOLD, MID
for col in range(2, 9):
    s.cell(tr, col).border = BORDER
s.cell(tr + 1, 2, "Ref. hours counts Lara's phases only. Lara's four phases total 6 weeks, 30 days, 180 hours at 50 m2.").font = GREY

nr = tr + 3
notes = [
    "How duration is calculated",
    "Working days = ref. weeks x days per week x (area / ref. area) ^ exponent, then x complexity / share of Lara's time (Lara phases only).",
    "Exponent 1.0 would be strictly linear. Below 1.0, a larger apartment takes less time per m2. 0 means fixed regardless of size.",
    "The exponents are my design judgement, not measured data. Compare against 3 or 4 finished projects and adjust.",
    "Complexity (Schedule sheet) is a per project multiplier: 1.0 normal, above 1.0 for heavy joinery or tricky layouts.",
    "Share of Lara's time: 100% means the reference pace. Two parallel projects at 50% each run at half pace, so each takes twice as long.",
]
for i, t in enumerate(notes):
    s.cell(nr + i, 2, t).font = BOLD if i == 0 else f(color="444444")

s["J3"] = "Days off (holidays, leave)"
s["J3"].font = BOLD
s["K3"] = "One date per cell. Skipped by all date formulas."
s["K3"].font = GREY
for r in range(4, 24):
    c = s.cell(r, 10)
    c.fill, c.font, c.number_format, c.border = INPUT_FILL, BLUE, "dd mmm yyyy", BORDER
HOL = "Settings!$J$4:$J$23"
for col, w in zip("ABCDEFGHIJK", (2, 34, 12, 18, 18, 18, 12, 78, 3, 18, 40)):
    s.column_dimensions[col].width = w

# ============================ Schedule ============================
sc = wb.create_sheet("Schedule")
sc.sheet_view.showGridLines = False
sc["A1"] = "Designature Studio project schedule"
sc["A1"].font = f(bold=True, size=14)
sc["A2"] = ("Fill the blue columns B to H and Notes. Each project has its own start date, so projects run in parallel. "
            "The Lara load row at the bottom shows when she is overbooked. Examples 01 to 06 are invented placeholders.")
sc["A2"].font = GREY

HR = 4
r0, rN = HR + 1, HR + MAX_ROWS
C_START, C_PH0 = 7, 9
def pc(k, j):   # phase k, j: 0 days, 1 start, 2 end
    return C_PH0 + 3 * k + j
C_FIN = C_PH0 + 3 * NP
C_HRS, C_NOTE = C_FIN + 1, C_FIN + 2
G0 = C_NOTE + 2

heads = ["#", "Project", "Client", "Area (m2)", "Complexity", "Share of Lara's time", "Start", "Status"]
for p in PHASES:
    heads += [f"{p[0]}\ndays", "start", "end"]
heads += ["Finish", "Lara hours", "Notes"]
for i, h in enumerate(heads, 1):
    c = sc.cell(HR, i, h)
    c.font, c.fill, c.alignment, c.border = WHITE_B, HEAD_FILL, CENTER, BORDER
sc.row_dimensions[HR].height = 54

# owner tag row above the phase headers
for k, p in enumerate(PHASES):
    c = sc.cell(HR - 1, pc(k, 0), p[1])
    c.font, c.alignment = f(size=8, italic=True, color="666666"), MID
    c.value = f"=Settings!$C${PH0 + k}"
    sc.merge_cells(start_row=HR - 1, start_column=pc(k, 0), end_row=HR - 1, end_column=pc(k, 2))

for n in range(MAX_ROWS):
    r = r0 + n
    p = PROJECTS[n] if n < len(PROJECTS) else (None,) * 8
    name, client, area, cx, share, start, status, note = p
    sc.cell(r, 1, n + 1).font = f(color="888888")
    for col, v in ((2, name), (3, client), (4, area), (5, cx), (6, share), (7, start), (8, status), (C_NOTE, note or None)):
        c = sc.cell(r, col, v)
        c.font, c.fill, c.border = BLUE, INPUT_FILL, BORDER
    for col in (4, 5, 6, 8):
        sc.cell(r, col).alignment = MID
    sc.cell(r, 5).number_format = "0.0"
    sc.cell(r, 6).number_format = "0%"
    sc.cell(r, 7).number_format = "dd mmm yy"

    for k in range(NP):
        cd, cs, ce = pc(k, 0), pc(k, 1), pc(k, 2)
        sr = PH0 + k
        D, S, E = f"{L(cd)}{r}", f"{L(cs)}{r}", f"{L(ce)}{r}"
        mult = f'IF(Settings!$C${sr}="Lara",$E{r}/$F{r},1)'
        sc.cell(r, cd, f'=IF($D{r}="",0,IF(Settings!$D${sr}=0,0,MAX(1,ROUND(Settings!$D${sr}*Settings!$C$5*($D{r}/Settings!$C$6)^Settings!$E${sr}*{mult},0))))')
        if k == 0:
            sc.cell(r, cs, f'=IF($G{r}="",0,WORKDAY($G{r}-1,1,{HOL}))')
        else:
            pe = f"{L(pc(k - 1, 2))}{r}"
            sc.cell(r, cs, f'=IF({pe}=0,0,IF({D}=0,{pe},WORKDAY({pe},1+Settings!$C$7,{HOL})))')
        sc.cell(r, ce, f'=IF({S}=0,0,IF({D}=0,{S},WORKDAY({S},{D}-1,{HOL})))')
        for col in (cd, cs, ce):
            sc.cell(r, col).font, sc.cell(r, col).border, sc.cell(r, col).alignment = BLACK, BORDER, MID
        sc.cell(r, cd).number_format = "0;;"
        sc.cell(r, cs).number_format = sc.cell(r, ce).number_format = DATE_FMT

    last = f"{L(pc(NP - 1, 2))}{r}"
    sc.cell(r, C_FIN, f"={last}").number_format = DATE_FMT
    terms = "+".join(f'{L(pc(k, 0))}{r}*(Settings!$C${PH0 + k}="Lara")' for k in range(NP))
    sc.cell(r, C_HRS, f'=IF($D{r}="",0,({terms})*$F{r}*Settings!$C$4)').number_format = "0;;"
    for col in (C_FIN, C_HRS):
        sc.cell(r, col).font, sc.cell(r, col).border, sc.cell(r, col).alignment = BLACK, BORDER, MID

# ---- Gantt ----
for w in range(GANTT_WEEKS):
    col = G0 + w
    c = sc.cell(HR, col, "=Settings!$C$8" if w == 0 else f"={L(col - 1)}{HR}+7")
    c.number_format = "dd mmm"
    c.font, c.fill = f(bold=True, size=8, color="FFFFFF"), HEAD_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
    sc.column_dimensions[L(col)].width = 4.6
    for r in range(r0, rN + 1):
        sc.cell(r, col).border = Border(left=Side(style="hair", color="DDDDDD"), right=Side(style="hair", color="DDDDDD"), top=thin, bottom=thin)
sc.cell(HR - 1, G0, "Weekly timeline, one column per week starting Monday").font = f(bold=True, size=9, color="666666")

g1, gN = L(G0), L(G0 + GANTT_WEEKS - 1)
rng = f"{g1}{r0}:{gN}{rN}"
for k, p in enumerate(PHASES):
    D, S, E = (f"${L(pc(k, j))}{r0}" for j in range(3))
    sc.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND({D}>0,{g1}${HR}<={E},{g1}${HR}+6>={S})"],
        fill=PatternFill("solid", bgColor=p[4], fgColor=p[4]), stopIfTrue=True))
sc.conditional_formatting.add(f"{g1}{HR}:{gN}{HR}", FormulaRule(
    formula=[f"AND(TODAY()>={g1}{HR},TODAY()<{g1}{HR}+7)"], fill=PatternFill("solid", bgColor="C0392B", fgColor="C0392B")))

# ---- Lara load row: share of her capacity booked each week ----
LR = rN + 1
sc.cell(LR, 2, "Lara load").font = BOLD
sc.cell(LR, 3, "Her phases only. Above 100% means overbooked.").font = f(size=8, italic=True, color="666666")
for w in range(GANTT_WEEKS):
    col = G0 + w
    wk = f"{L(col)}${HR}"
    parts = []
    for k in range(NP):
        S = f"${L(pc(k, 1))}${r0}:${L(pc(k, 1))}${rN}"
        E = f"${L(pc(k, 2))}${r0}:${L(pc(k, 2))}${rN}"
        D = f"${L(pc(k, 0))}${r0}:${L(pc(k, 0))}${rN}"
        sh = f"$F${r0}:$F${rN}"
        raw = f"({E}-({E}>{wk}+4)*({E}-{wk}-4)-{S}-({S}<{wk})*({wk}-{S})+1)"
        parts.append(f'SUMPRODUCT(({D}>0)*({raw}>0)*{raw}*{sh})*(Settings!$C${PH0 + k}="Lara")')
    c = sc.cell(LR, col, "=(" + "+".join(parts) + ")/Settings!$C$5")
    c.number_format, c.font, c.alignment = "0%;;", f(size=7), Alignment(horizontal="center")
    c.border = BORDER
sc.conditional_formatting.add(f"{g1}{LR}:{gN}{LR}", FormulaRule(
    formula=[f"{g1}{LR}>1.001"], fill=PatternFill("solid", bgColor="F4B6AE", fgColor="F4B6AE"), font=Font(name=FONT, size=7, bold=True, color="9C1C0F")))
sc.conditional_formatting.add(f"{g1}{LR}:{gN}{LR}", FormulaRule(
    formula=[f"AND({g1}{LR}>0,{g1}{LR}<=1.001)"], fill=PatternFill("solid", bgColor="D5E8D4", fgColor="D5E8D4")))

dv = DataValidation(type="list", formula1='"Planned,Active,On hold,Done"', allow_blank=True)
sc.add_data_validation(dv)
dv.add(f"H{r0}:H{rN}")

# ---- legend ----
lg = LR + 2
sc.cell(lg, 2, "Legend").font = BOLD
for k, p in enumerate(PHASES):
    c = sc.cell(lg + 1 + k, 2, p[0])
    c.fill, c.font = PatternFill("solid", fgColor=p[4]), f(color="000000")
    sc.cell(lg + 1 + k, 3, f"=Settings!$C${PH0 + k}").font = GREY
sc.cell(lg + NP + 2, 2, "Red header = current week.  Grey and sand bars are Anahit's tracking phases. Lara load counts only phases whose Owner is Lara on the Settings sheet.").font = GREY

widths = {1: 4, 2: 16, 3: 14, 4: 9, 5: 11, 6: 11, 7: 11, 8: 10, C_FIN: 11, C_HRS: 10, C_NOTE: 34}
for i in range(1, C_NOTE + 1):
    sc.column_dimensions[L(i)].width = widths.get(i, 9)
sc.column_dimensions[L(G0 - 1)].width = 2
sc.freeze_panes = sc.cell(HR + 1, 3)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
